import openpyxl
from openpyxl.utils import get_column_letter
from src.styles import StylesManager
from src.database_builder import DatabaseBuilder
from src.dashboard_builder import DashboardBuilder


class InvestmentSimulatorApp:
    def __init__(self, filename="investment_simulator.xlsx"):
        self.filename = filename
        self.wb = openpyxl.Workbook()
        self.styles = StylesManager()

    def run(self):
        # 1. Cria o banco de dados (Database_Profiles)
        db_builder = DatabaseBuilder(self.wb, self.styles)
        db_builder.build()

        # 2. Cria o Dashboard principal (APP)
        dash_builder = DashboardBuilder(self.wb, self.styles)
        dash_builder.build()

        # 3. Ajuste automático de largura de colunas
        for ws in self.wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        self.wb.save(self.filename)
        print(f"Sucesso absoluto! Simulador modular gerado em '{self.filename}'.")


if __name__ == "__main__":
    app = InvestmentSimulatorApp()
    app.run()
